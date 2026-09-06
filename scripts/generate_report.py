"""
Tạo file Word báo cáo Mini-Project 2 — VKU Room Booking App
"""

from docx import Document
from docx.shared import Pt, RGBColor, Inches, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import copy

# ── Màu sắc ──────────────────────────────────────────────────────────
PRIMARY   = RGBColor(0x1E, 0x6F, 0xD9)   # Xanh VKU
DARK      = RGBColor(0x11, 0x18, 0x27)   # Đen đậm
GRAY      = RGBColor(0x6B, 0x72, 0x80)   # Xám
WHITE     = RGBColor(0xFF, 0xFF, 0xFF)
LIGHT_BG  = RGBColor(0xF0, 0xF4, 0xFF)   # Xanh nhạt
GREEN     = RGBColor(0x06, 0x95, 0x55)

doc = Document()

# ── Page margins ─────────────────────────────────────────────────────
for section in doc.sections:
    section.top_margin    = Cm(2.0)
    section.bottom_margin = Cm(2.0)
    section.left_margin   = Cm(2.5)
    section.right_margin  = Cm(2.0)

def set_cell_bg(cell, hex_color):
    tc   = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd  = OxmlElement('w:shd')
    shd.set(qn('w:val'),   'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'),  hex_color)
    tcPr.append(shd)

def set_cell_border(cell, **kwargs):
    tc   = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = OxmlElement('w:tcBorders')
    for edge in ('top','left','bottom','right','insideH','insideV'):
        attrs = kwargs.get(edge, {})
        tag   = OxmlElement(f'w:{edge}')
        for k, v in attrs.items():
            tag.set(qn(f'w:{k}'), v)
        tcBorders.append(tag)
    tcPr.append(tcBorders)

def add_run(para, text, bold=False, size=11, color=None, italic=False):
    run = para.add_run(text)
    run.bold   = bold
    run.italic = italic
    run.font.size  = Pt(size)
    run.font.color.rgb = color if color else DARK
    return run

def heading(text, level=1, color=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after  = Pt(6)
    if level == 1:
        add_run(p, text, bold=True, size=16, color=color or PRIMARY)
    elif level == 2:
        add_run(p, text, bold=True, size=13, color=color or PRIMARY)
    else:
        add_run(p, text, bold=True, size=11, color=color or DARK)
    return p

def body(text, space_after=6, color=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    add_run(p, text, size=11, color=color or DARK)
    return p

def bullet(text, bold_prefix=None):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_after = Pt(3)
    if bold_prefix:
        add_run(p, bold_prefix, bold=True, size=11)
        add_run(p, text, size=11)
    else:
        add_run(p, text, size=11)

def hr():
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after  = Pt(4)
    pPr  = p._p.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    bot  = OxmlElement('w:bottom')
    bot.set(qn('w:val'),  'single')
    bot.set(qn('w:sz'),   '6')
    bot.set(qn('w:space'),'1')
    bot.set(qn('w:color'),'1E6FD9')
    pBdr.append(bot)
    pPr.append(pBdr)

# ═══════════════════════════════════════════════════════════════════════
# COVER
# ═══════════════════════════════════════════════════════════════════════
# Logo / tên trường (dùng text)
school = doc.add_paragraph()
school.alignment = WD_ALIGN_PARAGRAPH.CENTER
school.paragraph_format.space_after = Pt(2)
add_run(school, 'ĐẠI HỌC CÔNG NGHỆ VIỆT - HÀN (VKU)', bold=True, size=13, color=PRIMARY)

dept = doc.add_paragraph()
dept.alignment = WD_ALIGN_PARAGRAPH.CENTER
dept.paragraph_format.space_after = Pt(24)
add_run(dept, 'Khoa Công nghệ Thông tin', size=12, color=GRAY)

hr()

title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.paragraph_format.space_before = Pt(20)
title.paragraph_format.space_after  = Pt(6)
add_run(title, 'MINI-PROJECT SHORT TECHNICAL REPORT', bold=True, size=20, color=PRIMARY)

subtitle = doc.add_paragraph()
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
subtitle.paragraph_format.space_after = Pt(4)
add_run(subtitle, 'Mini-Project 2 — VKU Room Booking App', bold=True, size=14, color=DARK)

course_p = doc.add_paragraph()
course_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
course_p.paragraph_format.space_after = Pt(24)
add_run(course_p, 'Course: Cross-Platform Mobile App Development', size=12, color=GRAY, italic=True)

hr()

# Info table
info_table = doc.add_table(rows=4, cols=2)
info_table.alignment = WD_TABLE_ALIGNMENT.CENTER
info_table.style = 'Table Grid'
infos = [
    ('Sinh viên thực hiện:', 'Nguyễn Văn Bảo'),
    ('Mã số sinh viên:',    '21IT001'),
    ('Ngày nộp:',           '06/09/2026'),
    ('Giảng viên hướng dẫn:', '(Tên giảng viên)'),
]
for i, (label, val) in enumerate(infos):
    row = info_table.rows[i]
    row.cells[0].text = label
    row.cells[1].text = val
    for j, cell in enumerate(row.cells):
        for para in cell.paragraphs:
            for run in para.runs:
                run.bold = (j == 0)
                run.font.size = Pt(11)
                run.font.color.rgb = PRIMARY if j == 0 else DARK
        cell.width = Inches(2.5) if j == 0 else Inches(3.5)

doc.add_paragraph()
doc.add_page_break()

# ═══════════════════════════════════════════════════════════════════════
# SECTION 1
# ═══════════════════════════════════════════════════════════════════════
heading('1. GENERAL INFORMATION & DELIVERABLE LINKS')
hr()

body('Team Members:', color=PRIMARY)
bullet('Nguyễn Văn Bảo  —  Student ID: 21IT001  —  Role: Fullstack Developer  —  Contribution: 100%')

doc.add_paragraph()
links_table = doc.add_table(rows=3, cols=2)
links_table.style = 'Table Grid'
links_data = [
    ('🔗  Live Demo URL',    'https://vku-room-booking.vercel.app'),
    ('💻  GitHub Repository','https://github.com/nguyenvanbaoub2005/vku-room-booking'),
    ('🎥  Video Demo',       '(Chèn link video demo tại đây)'),
]
for i, (label, val) in enumerate(links_data):
    set_cell_bg(links_table.rows[i].cells[0], 'EFF6FF')
    links_table.rows[i].cells[0].text = label
    links_table.rows[i].cells[1].text = val
    for run in links_table.rows[i].cells[0].paragraphs[0].runs:
        run.bold = True
        run.font.color.rgb = PRIMARY
        run.font.size = Pt(10.5)
    for run in links_table.rows[i].cells[1].paragraphs[0].runs:
        run.font.size = Pt(10.5)

doc.add_paragraph()

# ═══════════════════════════════════════════════════════════════════════
# SECTION 2 — Feature Checklist
# ═══════════════════════════════════════════════════════════════════════
heading('2. FEATURE IMPLEMENTATION CHECKLIST')
hr()

feat_table = doc.add_table(rows=4, cols=4)
feat_table.style = 'Table Grid'
feat_table.alignment = WD_TABLE_ALIGNMENT.CENTER

headers = ['#', 'Required Feature', 'Status', 'Implementation Details']
for j, h in enumerate(headers):
    cell = feat_table.rows[0].cells[j]
    cell.text = h
    set_cell_bg(cell, '1E6FD9')
    for run in cell.paragraphs[0].runs:
        run.bold = True
        run.font.color.rgb = WHITE
        run.font.size = Pt(10.5)
    cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER

features = [
    ('1', 'Responsive\nMobile Viewport', '✅ Complete',
     '100% responsive đa nền tảng. Light Theme (trắng) cho Web Browser, Dark Theme (tối) cho iOS/Android. '
     'Hệ màu quản lý qua constants/Colors.ts với điều kiện Platform.OS === "web".'),
    ('2', 'Local Offline\nPersistence', '✅ Complete',
     'Zustand + AsyncStorage lưu Auth Token và Booking History cục bộ. '
     'Mọi thao tác đặt phòng ghi xuống @vku_bookings key ngay lập tức. '
     'Field synced: boolean đánh dấu trạng thái từng booking.'),
    ('3', 'Automatic\nBackground Sync', '✅ Complete',
     'Hook useNetworkMonitor theo dõi kết nối real-time (Web: window online/offline events; '
     'Native: expo-network + AppState). Khi kết nối khôi phục, syncPending() tự động chạy. '
     'NetworkBanner hiển thị: Offline → Đang đồng bộ → Thành công.'),
]

for i, (num, feat, status, detail) in enumerate(features):
    row = feat_table.rows[i + 1]
    row.cells[0].text = num
    row.cells[1].text = feat
    row.cells[2].text = status
    row.cells[3].text = detail
    bg = 'F8FAFF' if i % 2 == 0 else 'FFFFFF'
    for j, cell in enumerate(row.cells):
        set_cell_bg(cell, bg)
        for run in cell.paragraphs[0].runs:
            run.font.size = Pt(10)
            run.font.color.rgb = GREEN if j == 2 else DARK
            run.bold = (j == 2)
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER if j in (0, 2) else WD_ALIGN_PARAGRAPH.LEFT

doc.add_paragraph()

# ═══════════════════════════════════════════════════════════════════════
# SECTION 3 — Architecture
# ═══════════════════════════════════════════════════════════════════════
heading('3. TECHNICAL ARCHITECTURE & PROJECT STRUCTURE')
hr()

heading('3.1  Công nghệ sử dụng', level=2)

tech_table = doc.add_table(rows=7, cols=3)
tech_table.style = 'Table Grid'
tech_headers = ['Layer', 'Technology', 'Mục đích']
for j, h in enumerate(tech_headers):
    cell = tech_table.rows[0].cells[j]
    cell.text = h
    set_cell_bg(cell, '1E6FD9')
    for run in cell.paragraphs[0].runs:
        run.bold = True; run.font.color.rgb = WHITE; run.font.size = Pt(10.5)

tech_data = [
    ('Framework',    'React Native + Expo SDK 57', 'Cross-platform rendering (iOS, Android, Web)'),
    ('Routing',      'Expo Router (file-based)',    'Navigation & Deep linking'),
    ('State',        'Zustand v5',                  'Global state management'),
    ('Persistence',  'AsyncStorage',                'Offline local storage'),
    ('Network',      'expo-network + Browser Events','Network status detection'),
    ('Deployment',   'Vercel',                      'PWA hosting & CI/CD'),
]
for i, (layer, tech, desc) in enumerate(tech_data):
    row = tech_table.rows[i + 1]
    row.cells[0].text = layer; row.cells[1].text = tech; row.cells[2].text = desc
    bg = 'F0F4FF' if i % 2 == 0 else 'FFFFFF'
    for cell in row.cells:
        set_cell_bg(cell, bg)
        for run in cell.paragraphs[0].runs:
            run.font.size = Pt(10)

doc.add_paragraph()
heading('3.2  Luồng Background Sync', level=2)

flow_steps = [
    ('Người dùng đặt phòng (offline/online)',
     'Booking được tạo cục bộ → AsyncStorage.save() → pendingSync.push()'),
    ('Mất kết nối',
     'NetworkBanner đỏ: "🔴 Mất kết nối · N đặt phòng chờ đồng bộ"'),
    ('Kết nối được khôi phục',
     'useNetworkMonitor phát hiện → gọi onReconnect() → syncPending()'),
    ('Đang đồng bộ',
     'NetworkBanner vàng: "🟡 Đang đồng bộ N đặt phòng..."'),
    ('Hoàn tất',
     'Mark synced: true → NetworkBanner xanh: "🟢 Đã đồng bộ thành công!"'),
]

flow_table = doc.add_table(rows=len(flow_steps), cols=2)
flow_table.style = 'Table Grid'
bgs = ['EFF6FF', 'FEF9C3', 'DCFCE7', 'FEF9C3', 'DCFCE7']
for i, (step, desc) in enumerate(flow_steps):
    flow_table.rows[i].cells[0].text = f'Bước {i+1}: {step}'
    flow_table.rows[i].cells[1].text = desc
    set_cell_bg(flow_table.rows[i].cells[0], 'DBEAFE')
    set_cell_bg(flow_table.rows[i].cells[1], bgs[i])
    for cell in flow_table.rows[i].cells:
        for run in cell.paragraphs[0].runs:
            run.font.size = Pt(10)

doc.add_paragraph()
heading('3.3  Xử lý ngoại lệ', level=2)
bullet('Lỗi AsyncStorage', bold_prefix='Try/catch ở mọi thao tác persist: ')
bullet('Lỗi Network API', bold_prefix='getNetworkStateAsync() bọc try/catch: ')
bullet('Platform differences', bold_prefix='Platform.OS check tại runtime: ')

doc.add_paragraph()

# ═══════════════════════════════════════════════════════════════════════
# SECTION 4 — Screenshots
# ═══════════════════════════════════════════════════════════════════════
heading('4. EMPIRICAL EVIDENCE & SCREENSHOTS')
hr()

screenshots = [
    ('Screenshot 1', 'Trang chủ — Web (Light Theme) vs Mobile (Dark Theme)',
     'Chụp màn hình song song: giao diện nền trắng trên Browser và nền tối trên Expo Go/điện thoại.'),
    ('Screenshot 2', 'Offline Mode — NetworkBanner đỏ',
     'Tắt WiFi → chụp màn hình banner đỏ: "Mất kết nối · N đặt phòng chờ đồng bộ".'),
    ('Screenshot 3', 'Background Sync hoạt động',
     'Bật lại WiFi → chụp banner vàng "Đang đồng bộ..." chuyển sang xanh "Đã đồng bộ thành công!".'),
    ('Screenshot 4', 'PWA Install Guide Modal',
     'Tab Hồ sơ → "Cài đặt ứng dụng (PWA)" → chụp modal hướng dẫn 3 nền tảng.'),
]

for num, title_text, hint in screenshots:
    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_before = Pt(10)
    add_run(p_title, f'📸  {num}: {title_text}', bold=True, size=11, color=PRIMARY)

    # Placeholder box
    img_table = doc.add_table(rows=1, cols=1)
    img_table.style = 'Table Grid'
    img_cell = img_table.rows[0].cells[0]
    img_cell.height = Cm(7)
    set_cell_bg(img_cell, 'F8FAFF')
    ph = img_cell.paragraphs[0]
    ph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    ph.paragraph_format.space_before = Pt(30)
    add_run(ph, '[ Chèn ảnh chụp màn hình vào đây ]', size=11, color=GRAY, italic=True)

    hint_p = doc.add_paragraph()
    hint_p.paragraph_format.space_after = Pt(10)
    add_run(hint_p, f'💡  {hint}', size=10, color=GRAY, italic=True)

# ═══════════════════════════════════════════════════════════════════════
# SECTION 5 — Challenges
# ═══════════════════════════════════════════════════════════════════════
heading('5. TECHNICAL CHALLENGES & RESOLUTIONS')
hr()

# Challenge 1
heading('Thách thức 1: Tab Bar lỗi render trên Web', level=2)

prob1 = doc.add_paragraph()
add_run(prob1, 'Vấn đề: ', bold=True, color=RGBColor(0xDC, 0x26, 0x26))
add_run(prob1, 'Khi chạy Expo trên Web (react-native-web), thuộc tính boxShadow dạng CSS string '
        'không được chấp nhận bởi tabBarStyle prop, sinh ra visual artifacts. '
        'Chiều cao Tab Bar mobile (64px) quá lớn so với Web.', size=11)

prob2 = doc.add_paragraph()
add_run(prob2, 'Giải pháp: ', bold=True, color=GREEN)
add_run(prob2, 'Tách riêng style bằng Platform.OS === "web" tại runtime. '
        'Xóa boxShadow, thay bằng elevation: 0. '
        'Điều chỉnh height và padding riêng biệt cho từng nền tảng.', size=11)

code1 = doc.add_paragraph()
code1.paragraph_format.left_indent = Cm(1)
set_cell_bg  # not used here
add_run(code1,
    'tabBarStyle: {\n'
    '  height: isWeb ? 56 : 64,\n'
    '  paddingBottom: isWeb ? 4 : 8,\n'
    '  elevation: isWeb ? 0 : 8,\n'
    '}',
    size=9.5, color=GRAY)

heading('Thách thức 2: iOS Safari không hỗ trợ PWA Install Prompt', level=2)

prob3 = doc.add_paragraph()
add_run(prob3, 'Vấn đề: ', bold=True, color=RGBColor(0xDC, 0x26, 0x26))
add_run(prob3, 'Safari trên iOS không có beforeinstallprompt event (khác Chrome/Android). '
        'Người dùng không biết cách cài đặt PWA.', size=11)

prob4 = doc.add_paragraph()
add_run(prob4, 'Giải pháp: ', bold=True, color=GREEN)
add_run(prob4, 'Xây dựng 2 lớp hỗ trợ:\n'
        '(1) SafariInstallBanner tự động detect iOS Safari qua navigator.userAgent + '
        'window.matchMedia("(display-mode: standalone)"), hiện banner hướng dẫn.\n'
        '(2) InstallGuideModal trong tab Hồ sơ với hướng dẫn từng bước cho 3 nền tảng.', size=11)

doc.add_paragraph()
hr()

footer_p = doc.add_paragraph()
footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
add_run(footer_p, 'VKU Room Booking  ·  Khoa CNTT  ·  Đại học Công nghệ Việt - Hàn  ·  v1.0.0',
        size=9.5, color=GRAY, italic=True)

# ── Save ─────────────────────────────────────────────────────────────
output_path = '/Users/nguyenvan/vku-room-booking/REPORT_MiniProject2.docx'
doc.save(output_path)
print(f'✅ Đã tạo: {output_path}')
